#!/usr/bin/env python3
"""生成集成测试用的真实样本文件。
用法: python3 tools/make_fixtures.py <输出目录>
依赖: python-docx openpyxl python-pptx，以及 soffice（用于生成 doc/xls/ppt/wps/et/dps）。
"""
import os, subprocess, sys, shutil
out = sys.argv[1]
os.makedirs(out, exist_ok=True)

import docx
d = docx.Document()
d.add_paragraph("采购合同").style
d.add_paragraph("甲方：北京找找科技有限公司，乙方：张三。违约金为合同总额的百分之五。")
t = d.add_table(rows=1, cols=2)
t.rows[0].cells[0].text = "报销单号"
t.rows[0].cells[1].text = "BX-20250101"
d.sections[0].header.paragraphs[0].text = "页眉机密文字"
d.save(f"{out}/合同.docx")

import openpyxl
from openpyxl.comments import Comment
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "报销"
ws["A1"] = "姓名"; ws["B1"] = "金额"
ws["A2"] = "张三"; ws["B2"] = 1234.5
ws["C3"] = "=B2*2"
ws["A3"] = "差旅费"
ws["A2"].comment = Comment("请核对发票", "x")
ws2 = wb.create_sheet("第二页")
ws2["D4"] = "隐藏的关键词"
wb.save(f"{out}/报销.xlsx")

from pptx import Presentation
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "季度汇报"
s.placeholders[1].text = "营收增长百分之二十"
s.notes_slide.notes_text_frame.text = "讲稿备注：强调增长"
p.save(f"{out}/汇报.pptx")

# PDF：先写两页 docx，再用 LibreOffice 导出（内嵌字体、带 ToUnicode，与 Word/WPS 导出的 PDF 相近）
d2 = docx.Document()
d2.add_paragraph("第一页：不可抗力条款适用于本协议。")
d2.add_page_break()
d2.add_paragraph("第二页：Invoice number 20250707 已付款。")
d2.save(f"{out}/_pdfsrc.docx")
if shutil.which("soffice"):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out, f"{out}/_pdfsrc.docx"],
                   check=True, capture_output=True)
    shutil.move(f"{out}/_pdfsrc.pdf", f"{out}/说明.pdf")
os.remove(f"{out}/_pdfsrc.docx")

open(f"{out}/gbk.txt", "wb").write("这是一个GBK编码的文件，包含报销事项。\n第二行".encode("gbk"))
open(f"{out}/big5.txt", "wb").write("這是一個大五碼的檔案，包含報銷事項。我們的合同在這裡。".encode("big5"))
open(f"{out}/utf16.txt", "wb").write("﻿带BOM的UTF-16文件，内容包含差旅。".encode("utf-16-le"))
open(f"{out}/a.py", "w", encoding="utf8").write("# 计算税额\ndef tax(x):\n    return x * 0.13\n")

# 老格式 / WPS 格式：用 LibreOffice 转换
so = shutil.which("soffice")
if so:
    def conv(src, fmt, ext):
        subprocess.run([so, "--headless", "--convert-to", f"{fmt}", "--outdir", out + "/_c", f"{out}/{src}"],
                       check=True, capture_output=True)
        shutil.move(f"{out}/_c/{os.path.splitext(src)[0]}.{fmt.split(':')[0]}", f"{out}/{os.path.splitext(src)[0]}.{ext}")
    conv("合同.docx", "doc", "doc")
    conv("报销.xlsx", "xls", "xls")
    conv("汇报.pptx", "ppt", "ppt")
    # wps/et/dps 的文件结构与 doc/xls/ppt 相同，用于验证扩展名映射
    shutil.copy(f"{out}/合同.doc", f"{out}/合同2.wps")
    shutil.copy(f"{out}/报销.xls", f"{out}/报销2.et")
    shutil.copy(f"{out}/汇报.ppt", f"{out}/汇报2.dps")
    shutil.rmtree(out + "/_c", ignore_errors=True)
try:
    import py7zr, tempfile
    d7 = tempfile.mkdtemp()
    os.makedirs(d7 + "/合同", exist_ok=True)
    open(d7 + "/合同/说明.txt", "w", encoding="utf8").write("这里提到海獭这个词")
    open(d7 + "/顶层.txt", "w", encoding="utf8").write("顶层文件 内容")
    with py7zr.SevenZipFile(f"{out}/归档.7z", "w") as z:
        z.write(d7 + "/合同/说明.txt", "合同/说明.txt"); z.write(d7 + "/顶层.txt", "顶层.txt")
    with py7zr.SevenZipFile(f"{out}/加密.7z", "w", password="secret") as z:
        z.set_encrypted_header(True)
        z.write(d7 + "/顶层.txt", "顶层.txt")
except ImportError:
    pass
print("ok")
